from pathlib import Path
from hashlib import sha256
from zipfile import ZipFile
import json
import re
import shutil
import xml.etree.ElementTree as ET

import fitz
from docx import Document
from PIL import Image, ImageChops


ROOT = Path(__file__).resolve().parent
DOCX = ROOT / "manual_taller_de_ideas.docx"
PDF = ROOT / "manual_taller_de_ideas.pdf"
QA = ROOT / "qa"
shutil.copy2(QA / "render" / PDF.name, PDF)
doc = Document(DOCX)
pdf = fitz.open(PDF)
ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
checks = []


def add(name, passed, method, coverage, evidence):
    checks.append({"criterio": name, "estado": "PASS" if passed else "FAIL", "metodo": method, "cobertura": coverage, "evidencia": evidence})


def norm(text):
    return re.sub(r"\s+", " ", text).strip()


dimensions = [[round(p.rect.width * 25.4 / 72, 3), round(p.rect.height * 25.4 / 72, 3), p.rotation] for p in pdf]
add("Cuatro páginas A4 verticales", len(pdf) == 4 and all(abs(w - 210) < 0.1 and abs(h - 297) < 0.1 and r == 0 for w, h, r in dimensions), "PyMuPDF abre el PDF y mide MediaBox y rotación", "Las cuatro páginas del PDF final", dimensions)

groups = [[]]
for para in doc.paragraphs:
    if para.paragraph_format.page_break_before:
        groups.append([])
    if para.text.strip():
        groups[-1].append(para.text)
comparison = []
for i, page in enumerate(pdf):
    extracted = page.get_text(sort=True)
    extracted = re.sub(r"Taller de ideas\s*\d\s*/\s*4\s*$", "", extracted)
    comparison.append(norm(extracted) == norm("\n".join(groups[i])))
add("Contenido íntegro y en la página correspondiente", len(groups) == 4 and all(comparison), "Comparación de todo el texto nativo del DOCX por etapa con la extracción del PDF; normalización de espacios y exclusión del pie", "Todo el texto del cuerpo, incluidas etiquetas y títulos", {"coincidencia_por_pagina": comparison, "caracteres_docx_sin_espacios_de_formulario": sum(len(norm(" ".join(g))) for g in groups)})

embedded = []
for xref in sorted({font[0] for page in pdf for font in page.get_fonts()}):
    fontname, ext, fonttype, fontdata = pdf.extract_font(xref)
    embedded.append({"xref": xref, "nombre": fontname, "extension": ext, "bytes_incrustados": len(fontdata)})
add("Fuentes incluidas en PDF", bool(embedded) and all(f["bytes_incrustados"] > 0 for f in embedded), "Extracción de los programas de fuente por referencia PDF con PyMuPDF", "Todas las fuentes referenciadas en las cuatro páginas", embedded)

grayscale = []
for index in range(1, 5):
    im = Image.open(QA / "render" / f"page-{index}.png").convert("RGB")
    red, green, blue = im.split()
    max_diff = max(ImageChops.difference(red, green).getextrema()[1], ImageChops.difference(green, blue).getextrema()[1])
    grayscale.append({"pagina": index, "diferencia_maxima_entre_canales": max_diff, "pixeles": list(im.size)})
add("Render en escala de grises", all(g["diferencia_maxima_entre_canales"] == 0 for g in grayscale), "Comparación RGB de todos los píxeles de las imágenes renderizadas por Poppler a 135 ppp", "Las cuatro imágenes finales; no es prueba de impresora", grayscale)

margins = []
for index, page in enumerate(pdf):
    rects = [fitz.Rect(span["bbox"]) for block in page.get_text("dict")["blocks"] if "lines" in block for line in block["lines"] for span in line["spans"] if span["text"].strip()]
    rects.extend(d["rect"] for d in page.get_drawings())
    left = min(r.x0 for r in rects)
    top = min(r.y0 for r in rects)
    right = page.rect.width - max(r.x1 for r in rects)
    bottom = page.rect.height - max(r.y1 for r in rects)
    margins.append({"pagina": index + 1, "margenes_minimos_mm": [round(v * 25.4 / 72, 2) for v in (left, top, right, bottom)]})
add("Área imprimible prevista", all(min(m["margenes_minimos_mm"]) >= 8 for m in margins), "Medición de las cajas de texto y trazos del PDF; mínimo elegido de 8 mm en cada borde", "Contenido y pies de las cuatro páginas, para una comprobación digital de geometría", margins)

with ZipFile(DOCX) as package:
    document_xml = ET.fromstring(package.read("word/document.xml"))
    settings_xml = ET.fromstring(package.read("word/settings.xml"))
    media = [n for n in package.namelist() if n.startswith("word/media/")]
    drawings = document_xml.findall(".//w:drawing", ns)
    protections = settings_xml.findall(".//w:documentProtection", ns)
    native_paragraphs = document_xml.findall(".//w:body/w:p", ns)
    native_runs = document_xml.findall(".//w:t", ns)
add("Fuente DOCX con contenido nativo", not media and not drawings and not protections and len(native_runs) > 50, "Inspección OOXML del paquete DOCX y apertura con python-docx", "Texto, títulos, campos de escritura, reglas y estructura del documento", {"parrafos": len(native_paragraphs), "nodos_texto": len(native_runs), "archivos_media": len(media), "dibujos": len(drawings), "protecciones": len(protections), "estilos_usados": sorted({p.style.name for p in doc.paragraphs})})

notes_lines = []
last = pdf[-1]
notes_span = next(span for b in last.get_text("dict")["blocks"] if "lines" in b for line in b["lines"] for span in line["spans"] if span["text"].strip() == "Notas")
for drawing in last.get_drawings():
    r = drawing["rect"]
    if r.y0 > notes_span["bbox"][3] and r.width > 400:
        notes_lines.append(r)
raw_note_traces = len(notes_lines)
notes_lines = sorted([fitz.Rect(c) for c in {tuple(round(v, 4) for v in r) for r in notes_lines}], key=lambda r: r.y0)
add("Espacio para notas manuscritas", len(notes_lines) == 7, "Conteo y medición de los trazos bajo el rótulo Notas en el PDF; agrupación de los trazos exactamente coincidentes", "Área de notas de la página 4", {"lineas_distintas": len(notes_lines), "trazos_pdf_brutos": raw_note_traces, "ancho_mm": round(notes_lines[0].width * 25.4 / 72, 2), "distancias_entre_lineas_mm": [round((b.y0 - a.y0) * 25.4 / 72, 2) for a, b in zip(notes_lines, notes_lines[1:])]})

digest = sha256(DOCX.read_bytes()).hexdigest()
edited = Document(DOCX)
target = next(p for p in edited.paragraphs if p.text == "1  Acuerden una pregunta")
target.runs[0].text = "1  Acuerden una pregunta clara"
index = next(i for i, p in enumerate(edited.paragraphs) if p.text == "Nuestra pregunta")
edited.paragraphs[index + 1].runs[0].text = "¿Cómo podemos organizar mejor nuestras reuniones?"
edit_dir = QA / "editabilidad"
edit_dir.mkdir(exist_ok=True)
edited_path = edit_dir / "prueba_edicion.docx"
edited.save(edited_path)
reopened = Document(edited_path)
new_text = "\n".join(p.text for p in reopened.paragraphs)
edit_passed = "1  Acuerden una pregunta clara" in new_text and "¿Cómo podemos organizar mejor nuestras reuniones?" in new_text and sha256(DOCX.read_bytes()).hexdigest() == digest
add("Edición y guardado de una copia DOCX", edit_passed, "Edición de un título de paso y un espacio de respuesta con python-docx; guardado y reapertura; hash del original antes y después", "Dos ediciones representativas; no es interacción en Microsoft Word", {"copia": str(edited_path.relative_to(ROOT)), "original_sha256_inalterado": digest})

report = {"encargo": "Manual A4 de cuatro páginas para taller de ideas", "productores": {"docx": "python-docx 1.2.0", "pdf": "LibreOfficeDev 26.8.0.0.alpha0 mediante render_docx.py", "raster": "Poppler a 135 ppp"}, "artefactos": {"docx": {"ruta": DOCX.name, "bytes": DOCX.stat().st_size, "sha256": digest}, "pdf": {"ruta": PDF.name, "bytes": PDF.stat().st_size, "sha256": sha256(PDF.read_bytes()).hexdigest()}}, "comprobaciones": checks}
(QA / "verificacion_digital.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(report, ensure_ascii=False, indent=2))
if not all(c["estado"] == "PASS" for c in checks):
    raise SystemExit(1)

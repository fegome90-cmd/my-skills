from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "manual_taller_de_ideas.docx"
FONT = "DejaVu Sans"
doc = Document()
section = doc.sections[0]
section.page_width = Mm(210)
section.page_height = Mm(297)
section.top_margin = Mm(18)
section.bottom_margin = Mm(18)
section.left_margin = Mm(20)
section.right_margin = Mm(20)
section.header_distance = Mm(8)
section.footer_distance = Mm(9)


def configure_style(name, size, bold=False, before=0, after=0, line=1.14):
    style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + attr), FONT)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + attr), None)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "es-ES")
    rpr.append(lang)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line
    style.paragraph_format.widow_control = True
    style.paragraph_format.keep_together = True
    for borders in list(style.element.iter(qn("w:pBdr"))):
        borders.getparent().remove(borders)
    return style


configure_style("Normal", 11.5, after=6)
configure_style("Title", 25, bold=True, after=10, line=1.08)
configure_style("Heading 1", 24, bold=True, after=10, line=1.08)
configure_style("Heading 2", 12.5, bold=True, before=11, after=3, line=1.08)
configure_style("Lead", 11.5, after=12, line=1.18)
configure_style("Stage", 9, bold=True, after=9, line=1)
configure_style("Field Label", 10.5, bold=True, before=10, after=0, line=1.1)
configure_style("Footer", 9, after=0, line=1)
for style in ("Title", "Heading 1", "Heading 2", "Stage", "Field Label"):
    doc.styles[style].paragraph_format.keep_with_next = True
doc.styles["Heading 1"].paragraph_format.page_break_before = False
doc.styles["Heading 2"].paragraph_format.page_break_before = False

core = doc.core_properties
core.title = "Organizar un pequeño taller de ideas"
core.subject = "Manual de trabajo en cuatro etapas"
core.author = ""
core.last_modified_by = ""
core.keywords = "taller, ideas, manual"
core.comments = ""
core.language = "es-ES"

settings = doc.settings.element
update = OxmlElement("w:updateFields")
update.set(qn("w:val"), "true")
settings.append(update)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1" if instruction == "PAGE" else "4"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for item in (begin, text, separate, value, end):
        run._r.append(item)


footer = section.footer.paragraphs[0]
footer.style = doc.styles["Footer"]
doc.styles["Footer"].paragraph_format.tab_stops.clear_all()
footer.paragraph_format.tab_stops.clear_all()
footer.paragraph_format.tab_stops.add_tab_stop(Mm(170), WD_TAB_ALIGNMENT.RIGHT)
footer.add_run("Taller de ideas\t")
add_field(footer, "PAGE")
footer.add_run(" / ")
add_field(footer, "NUMPAGES")


def p(text, style=None):
    return doc.add_paragraph(text, style)


def stage(number, title, first=False):
    kicker = p(f"ETAPA {number} DE 4", "Stage")
    if not first:
        kicker.paragraph_format.page_break_before = True
    p(title, "Title" if first else "Heading 1")


def step(number, title, text):
    p(f"{number}  {title}", "Heading 2")
    p(text)


def lines(label, count=1, height=21):
    if label:
        p(label, "Field Label")
    for _ in range(count):
        para = p("\u00a0")
        pf = para.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(height)
        pf.keep_with_next = False
        borders = OxmlElement("w:pBdr")
        for name in ("bottom", "between"):
            border = OxmlElement("w:" + name)
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "4")
            border.set(qn("w:color"), "9B9B9B")
            border.set(qn("w:space"), "0")
            borders.append(border)
        para._p.get_or_add_pPr().append(borders)


stage(1, "Organizar un pequeño\ntaller de ideas", first=True)
p("Una guía para pasar de una pregunta compartida a una idea que puedan probar. Usen estas páginas durante el taller y completen los espacios a mano.", "Lead")
p("Objetivo y preparación", "Heading 2")
p("El objetivo es elegir una idea para probar y acordar un siguiente paso con una persona responsable y una fecha.")
step(1, "Acuerden una pregunta", "Definan qué quieren mejorar o resolver. Escríbanlo como una pregunta concreta que todo el grupo entienda.")
lines("Nuestra pregunta", 2)
step(2, "Reúnan los materiales", "Preparen papel y lápices para cada persona. Dejen hojas adicionales a mano para agrupar ideas y dibujar una primera versión.")
step(3, "Preparen el espacio", "Busquen una superficie donde puedan trabajar y ver las ideas. Coloquen los materiales al alcance de todas las personas.")
step(4, "Compartan el propósito", "Expliquen qué quieren conseguir y anótenlo abajo. Mantengan ese propósito a la vista y confirmen que todas las personas entienden la pregunta.")
lines("Queremos conseguir", 2)

stage(2, "Explorar")
p("Hagan visibles las ideas. Escuchen cada propuesta antes de buscar relaciones entre ellas.", "Lead")
step(1, "Anoten por separado", "Cada persona escribe sus ideas. Pongan una idea en cada hoja o en cada bloque, con palabras sencillas o un dibujo.")
step(2, "Compartan sin interrumpir", "Por turnos, cada persona presenta lo que anotó. Dejen que termine; después hagan preguntas para entender. Reserven la valoración para la siguiente etapa.")
step(3, "Agrupen temas comunes", "Coloquen juntas las ideas que se parecen. Pongan un nombre sencillo a cada grupo y conserven también las ideas que no encajan aún.")
lines("Ideas que queremos conservar", 5, 24)
lines("Temas que aparecen", 4, 24)
p("Antes de seguir", "Heading 2")
p("Revisen los grupos y comprueben que todas las personas pudieron aportar.")

stage(3, "Elegir y probar")
p("Usen el propósito para decidir qué idea merece una primera prueba.", "Lead")
step(1, "Comparen con el propósito", "Vuelvan a la pregunta y al propósito de la primera página. Conversen sobre qué aporta cada idea y qué haría falta para probarla.")
step(2, "Escojan una idea", "Acuerden cuál van a desarrollar ahora. Dejen anotadas las demás para poder retomarlas.")
step(3, "Creen un prototipo sencillo", "Hagan una primera versión que permita explicar o ensayar la idea: un boceto, una secuencia de pasos o una representación breve.")
step(4, "Prueben y anoten ajustes", "Usen el prototipo para recorrer la idea. Registren qué se entiende, qué falta y qué conviene cambiar.")
lines("Idea elegida y por qué", 2, 23)
lines("Nuestro primer prototipo", 3, 23)
lines("Qué conviene ajustar", 2, 23)

stage(4, "Cerrar")
p("Cierren el encuentro con lo aprendido y un compromiso que quede por escrito.", "Lead")
step(1, "Compartan lo aprendido", "Digan qué descubrieron al explorar y probar. Anoten las dudas que siguen abiertas.")
step(2, "Acuerden el siguiente paso", "Definan una acción concreta. Registren una persona responsable y una fecha acordada.")
lines("Lo que aprendimos", 2, 22)
lines("Siguiente paso", 1, 22)
lines("Persona responsable", 1, 22)
lines("Fecha acordada", 1, 22)
closing = p("Antes de terminar, lean el acuerdo en voz alta y comprueben que todas las personas lo entienden.")
closing.paragraph_format.space_before = Pt(10)
p("Notas", "Heading 2")
lines(None, 7, 24)

doc.save(OUT)
print(OUT)
